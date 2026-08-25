"""Apply a structured tailoring patch to a ``.docx``, deterministically.

prd.md (Resume Revisions): the model never edits a resume template itself. It
sees each template as numbered paragraphs and returns a JSON patch that names
the template it targets (``base``) and lists per-paragraph operations —
``replace`` a paragraph's text, ``insert_after`` a new paragraph (which
inherits the anchor paragraph's formatting, so a bullet inserted after a
bullet is itself a bullet), ``delete`` one, or ``move`` one to sit ``after`` or
``before`` another paragraph (an explicit reorder op, so leading a role with
its most relevant bullet is a mechanical relocation that preserves the run
verbatim, not a delete-and-retype that risks dropping bold). This module
applies that patch to an in-memory copy of the chosen template with
python-docx. Every id a change targets is resolved against the *original*
numbering before anything moves, so inserts, deletes, and moves never shift the
ids still to be applied, and bold
spans ride inline as ``**double asterisks**`` rendered into runs. The split is
what buys a changelog rendered *from the applied patch* rather than a second
model artifact that could drift, and formatting inherited from the template's
own runs rather than synthesized.

Everything here is pure logic over python-docx's in-memory ``Document``
objects — no network, no database, no filesystem — so it is directly testable
(per the repo convention of keeping pure logic I/O-free).
"""

from __future__ import annotations

import json
import re
from copy import deepcopy
from dataclasses import dataclass
from typing import Literal

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, Field, ValidationError, model_validator

from .search.parse import extract_json_text


class PatchError(ValueError):
    """The model's patch cannot be applied to the base resume as returned."""


class PatchChange(BaseModel):
    """One paragraph operation, as the tailoring model returns it.

    ``op`` is ``replace`` (swap the paragraph's text), ``insert_after`` (add a
    new paragraph after this one, inheriting its formatting), ``delete``
    (remove it), or ``move`` (relocate it to sit immediately ``after`` or
    ``before`` another paragraph). ``paragraph`` is always a ``[P<n>]`` id from
    the *original* numbering; ``after``/``before`` (exactly one, on a move) are
    likewise original ids naming the anchor to land next to; ``text`` is the
    full paragraph text (with ``**bold**`` spans) for replace/insert and is
    unused for delete/move.
    """

    op: Literal["replace", "insert_after", "delete", "move"] = "replace"
    paragraph: int
    after: int | None = None
    before: int | None = None
    text: str = ""
    rationale: str

    @model_validator(mode="after")
    def _validate_shape(self) -> PatchChange:
        if self.op in ("replace", "insert_after") and not self.text.strip():
            raise ValueError(
                f"a {self.op} change on paragraph {self.paragraph} needs non-empty text"
            )
        if self.op == "move":
            anchors = [a for a in (self.after, self.before) if a is not None]
            if len(anchors) != 1:
                raise ValueError(
                    f"a move change on paragraph {self.paragraph} needs exactly one of "
                    "'after' or 'before'"
                )
            if anchors[0] == self.paragraph:
                raise ValueError(
                    f"a move change cannot anchor paragraph {self.paragraph} on itself"
                )
        elif self.after is not None or self.before is not None:
            raise ValueError(
                f"'after'/'before' are only valid on a move change (paragraph {self.paragraph})"
            )
        return self


class TailoringPatch(BaseModel):
    """The model's whole answer: the template pick plus the change list.

    ``base`` names the resume template the patch targets (required — the model
    picks the template, per prd.md decided 2026-08-22) and ``base_rationale``
    is why, rendered into the changelog the user reviews. ``new_family``,
    normally null, declares that no template's role family fits: the tailored
    result then seeds the library's new family template (the outward-expansion
    rule). An empty ``changes`` list is a valid answer — the prompt explicitly
    allows \"no change is clearly better\" — and yields an unmodified copy of
    the chosen template plus a changelog saying so.
    """

    base: str
    base_rationale: str | None = None
    new_family: str | None = None
    summary: str | None = None
    changes: list[PatchChange] = Field(default_factory=list)


def parse_patch(raw: str) -> TailoringPatch:
    """Parse and validate the model's raw output into a ``TailoringPatch``.

    Tolerates markdown fences and surrounding prose (the same stance as the
    search-output parser); anything structurally wrong raises rather than
    letting a malformed patch half-apply.
    """
    try:
        data = json.loads(extract_json_text(raw))
    except (ValueError, json.JSONDecodeError) as exc:
        raise PatchError(f"the tailoring output is not valid JSON: {exc}") from exc
    try:
        return TailoringPatch.model_validate(data)
    except ValidationError as exc:
        raise PatchError(f"the tailoring output does not match the patch contract: {exc}") from exc


def iter_paragraphs(document):
    """Every paragraph in the document, body first, then table cells.

    Resumes are frequently laid out in tables, which ``document.paragraphs``
    does not descend into — so table-cell paragraphs are walked too. Merged
    cells repeat the same underlying element across a row; deduplicate on the
    element's identity so a paragraph is numbered exactly once. (Tables nested
    inside cells are not descended into — not a layout resumes use.)
    """
    yield from document.paragraphs
    seen: set[int] = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                key = id(cell._tc)
                if key in seen:
                    continue
                seen.add(key)
                yield from cell.paragraphs


def render_numbered_text(document) -> str:
    """The document as ``[P<n>] <text>`` lines — what the model sees.

    Empty paragraphs (spacers) are not rendered but keep their index, so the
    ids the model can target are exactly the ids shown, and they map straight
    back onto ``iter_paragraphs`` order in ``apply_patch``.
    """
    lines: list[str] = []
    for i, paragraph in enumerate(iter_paragraphs(document)):
        text = paragraph.text.strip()
        if text:
            lines.append(f"[P{i}] {text}")
    return "\n".join(lines)


@dataclass(frozen=True)
class AppliedChange:
    """One change as actually applied — the changelog's raw material.

    For a ``move``, ``before`` carries the relocated paragraph's text and
    ``anchor``/``relation`` name where it landed (``before``/``after`` which
    id); ``after`` is unused. The two trailing fields default so the
    replace/insert/delete constructions stay positional.
    """

    op: str
    paragraph: int
    before: str
    after: str
    rationale: str
    anchor: int | None = None
    relation: str = ""


_BOLD = re.compile(r"\*\*(.+?)\*\*", re.DOTALL)


def _bold_segments(text: str) -> list[tuple[str, bool]]:
    """Split ``text`` into ``(chunk, is_bold)`` segments on ``**...**`` markers."""
    segments: list[tuple[str, bool]] = []
    pos = 0
    for match in _BOLD.finditer(text):
        if match.start() > pos:
            segments.append((text[pos : match.start()], False))
        segments.append((match.group(1), True))
        pos = match.end()
    if pos < len(text):
        segments.append((text[pos:], False))
    return segments or [(text, False)]


def _set_runs(paragraph, text: str) -> None:
    """Rewrite a paragraph's runs to render ``text`` with ``**bold**`` spans.

    Character formatting (font, size) is inherited from the paragraph's first
    existing run, so a rewritten bullet keeps the template's typeface; bold is
    then set explicitly per segment, so the template's lead-phrase bold neither
    bleeds onto plain segments nor is lost on the ones the model marks bold. A
    paragraph with no runs (rare) just gets fresh unformatted runs.
    """
    existing = paragraph.runs
    template_rpr = None
    if existing:
        found = existing[0]._r.find(qn("w:rPr"))
        if found is not None:
            template_rpr = deepcopy(found)
    for run in list(existing):
        run._r.getparent().remove(run._r)
    for chunk, is_bold in _bold_segments(text):
        run = paragraph.add_run(chunk)
        if template_rpr is not None:
            run._r.insert(0, deepcopy(template_rpr))
        run.bold = is_bold


def _insert_paragraph_after(anchor):
    """Clone ``anchor`` and splice the copy in right after it; return the copy.

    Deep-copying the whole ``<w:p>`` carries the anchor's paragraph properties
    (style, list/numbering, spacing) and runs forward — the same way Word makes
    a new bullet inherit the previous bullet — so ``_set_runs`` then only has to
    swap the text.
    """
    new_p = deepcopy(anchor._p)
    anchor._p.addnext(new_p)
    return Paragraph(new_p, anchor._parent)


def apply_patch(document, patch: TailoringPatch) -> list[AppliedChange]:
    """Apply ``patch`` to ``document`` in place; return what actually changed.

    Anchors are resolved against the original numbering, then applied in
    phases (replace, insert, move, delete) so ids never shift mid-patch. A change
    targeting an id that was never offered — out of range, or an empty spacer
    paragraph — raises ``PatchError`` rather than guessing. A ``replace`` whose
    text restates the paragraph unchanged is dropped silently (the contract
    forbids no-ops, but one is not worth failing a whole resume over).
    """
    originals = list(iter_paragraphs(document))
    count = len(originals)

    # Resolve and check every anchor against the ORIGINAL numbering first, so
    # the ids the model targeted stay valid no matter how many inserts and
    # deletes follow.
    def _require_offered(idx: int) -> None:
        if not 0 <= idx < count:
            raise PatchError(
                f"the patch targets paragraph {idx}, but the base resume "
                f"has paragraphs 0–{count - 1}"
            )
        if not originals[idx].text.strip():
            raise PatchError(
                f"the patch targets empty paragraph {idx}, which was never "
                "offered in the numbered base text"
            )

    for change in patch.changes:
        _require_offered(change.paragraph)
        if change.op == "move":
            _require_offered(change.after if change.after is not None else change.before)

    applied: list[AppliedChange] = []
    # Phase 1 — replaces, mutating the original paragraphs in place.
    for change in patch.changes:
        if change.op != "replace":
            continue
        target = originals[change.paragraph]
        before = target.text
        if change.text == before:  # the contract forbids no-ops; drop, don't fail
            continue
        _set_runs(target, change.text)
        applied.append(
            AppliedChange("replace", change.paragraph, before, change.text, change.rationale)
        )
    # Phase 2 — inserts. New paragraphs are cloned siblings, so original ids
    # stay valid; multiple inserts after one anchor chain off the last insert
    # to keep the model's listed order.
    last_after: dict[int, object] = {}
    for change in patch.changes:
        if change.op != "insert_after":
            continue
        anchor = last_after.get(change.paragraph, originals[change.paragraph])
        new_para = _insert_paragraph_after(anchor)
        _set_runs(new_para, change.text)
        last_after[change.paragraph] = new_para
        applied.append(
            AppliedChange("insert_after", change.paragraph, "", change.text, change.rationale)
        )
    # Phase 3 — moves. Relocate an original paragraph to sit next to an anchor,
    # both resolved from the original numbering. lxml's addnext/addprevious
    # re-parents the element in place, so there is no clone and no renumber, and
    # a replace already applied to it in phase 1 rides along.
    for change in patch.changes:
        if change.op != "move":
            continue
        moved = originals[change.paragraph]
        if change.after is not None:
            originals[change.after]._p.addnext(moved._p)
            anchor, relation = change.after, "after"
        else:
            originals[change.before]._p.addprevious(moved._p)
            anchor, relation = change.before, "before"
        applied.append(
            AppliedChange(
                "move", change.paragraph, moved.text, "", change.rationale, anchor, relation
            )
        )
    # Phase 4 — deletes last, so an insert_after a to-be-deleted anchor still
    # lands (as a following sibling) before the anchor element is removed.
    for change in patch.changes:
        if change.op != "delete":
            continue
        target = originals[change.paragraph]
        before = target.text
        target._p.getparent().remove(target._p)
        applied.append(AppliedChange("delete", change.paragraph, before, "", change.rationale))
    return applied


def render_changelog(
    *,
    company: str,
    title: str,
    model: str,
    date_generated: str,
    base: str,
    base_rationale: str | None,
    new_template: str | None,
    summary: str | None,
    applied: list[AppliedChange],
) -> str:
    """Render ``resume_changelog.md`` from the applied patch. Pure.

    One addressable entry per change, each carrying its rationale — the
    artifact that lets an interactive review accept/reject change #N instead
    of re-deriving the whole diff by hand. Also records which template was
    chosen and why, and flags a new-template creation for the curation scrub
    prd.md requires.
    """
    lines = [
        f"# Resume changelog — {company} — {title}",
        "",
        f"Tailored {date_generated} by `jsa generate` ({model}, structured patch "
        f"applied to the `{base}` template).",
        "",
    ]
    if base_rationale:
        lines += [f"**Template choice:** {base_rationale}", ""]
    if new_template:
        lines += [
            f"> **NEW TEMPLATE CREATED:** this tailoring seeded "
            f"`resume_templates/{new_template}.docx` as a new role-family template. "
            "It began life tailored to this one posting — review it and scrub "
            "company-specific phrasing before its next use.",
            "",
        ]
    if summary:
        lines += [summary, ""]
    if not applied:
        lines += [
            "The tailoring pass proposed no changes — the base resume was emitted unchanged.",
            "",
        ]
        return "\n".join(lines)
    lines += [f"## Changes ({len(applied)})", ""]
    headings = {
        "replace": "Replaced paragraph",
        "insert_after": "Inserted after paragraph",
        "delete": "Deleted paragraph",
        "move": "Moved paragraph",
    }
    for n, change in enumerate(applied, start=1):
        heading = headings.get(change.op, "Paragraph")
        lines += [f"### {n}. {heading} {change.paragraph}", ""]
        if change.op == "move":
            lines.append(f"- **Moved:** {change.before}")
            lines.append(f"- **To:** {change.relation} paragraph {change.anchor}")
        else:
            if change.before:
                lines.append(f"- **Before:** {change.before}")
            if change.after:
                lines.append(f"- **After:** {change.after}")
        lines += [f"- **Why:** {change.rationale}", ""]
    return "\n".join(lines)
